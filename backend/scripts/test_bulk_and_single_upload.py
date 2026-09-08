import os
import sys
import io

# Ensure backend path is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../..")))

from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.core.database import SessionLocal
from backend.app.core.security import get_password_hash, create_access_token
from backend.app.models import User, UserRole, Student, Section, AuditLog

import pandas as pd
import docx

client = TestClient(app)

def run_tests():
    db = SessionLocal()
    print("=== Starting Comprehensive Upload Verification Tests ===")

    # Setup Admin User
    admin = db.query(User).filter(User.email == "test_admin@placematch.edu").first()
    if not admin:
        admin = User(
            email="test_admin@placematch.edu",
            hashed_password=get_password_hash("password123"),
            full_name="Test Admin",
            role=UserRole.ADMIN,
            is_active=True
        )
        db.add(admin)
        db.commit()
        db.refresh(admin)

    admin_token = create_access_token({"sub": admin.email, "role": admin.role.value})
    admin_headers = {"Authorization": f"Bearer {admin_token}"}

    # Setup Test Section
    sec = db.query(Section).filter(Section.name == "MultiFormat-UploadTest").first()
    if not sec:
        sec = Section(name="MultiFormat-UploadTest", created_by=admin.id)
        db.add(sec)
        db.commit()
        db.refresh(sec)
    section_id = sec.id

    # -------------------------------------------------------------------------
    # TEST 1: CORS Preflight Check
    # -------------------------------------------------------------------------
    print("\n--- TEST 1: CORS Preflight Check ---")
    cors_res = client.options(
        f"/api/v1/sections/{section_id}/upload-resume",
        headers={
            "Origin": "http://localhost:5173",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "authorization",
        }
    )
    print(f"CORS Options Status Code: {cors_res.status_code}")
    print(f"Access-Control-Allow-Origin: {cors_res.headers.get('access-control-allow-origin')}")
    assert cors_res.status_code == 200
    assert cors_res.headers.get("access-control-allow-origin") == "http://localhost:5173"
    print("[SUCCESS] CORS Preflight succeeded for http://localhost:5173!")

    # -------------------------------------------------------------------------
    # TEST 2: Single Student Resume PDF Upload
    # -------------------------------------------------------------------------
    print("\n--- TEST 2: Single Student Resume PDF Upload ---")
    # Generate minimal valid PDF using pypdf
    import pypdf
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    pdf_buffer = io.BytesIO()
    writer.write(pdf_buffer)
    pdf_bytes = pdf_buffer.getvalue()

    files = {"file": ("RESUME.pdf", pdf_bytes, "application/pdf")}
    params = {
        "email": "pdf_single_student@placematch.edu",
        "full_name": "PDF Single Student",
        "roll_number": "ROLL_PDF_01",
        "department": "Computer Science",
        "cgpa": 8.75,
        "active_backlogs": 0,
        "skills_csv": "Python, FastAPI, PyPDF",
    }

    res = client.post(f"/api/v1/sections/{section_id}/upload-resume", params=params, files=files, headers=admin_headers)
    print(f"Single PDF Upload Status: {res.status_code}")
    print(f"Response: {res.json()}")
    assert res.status_code == 200
    assert res.json()["roll_number"] == "ROLL_PDF_01"
    print("[SUCCESS] Single Student PDF Resume Upload succeeded!")

    # -------------------------------------------------------------------------
    # TEST 3: CSV Bulk Upload
    # -------------------------------------------------------------------------
    print("\n--- TEST 3: CSV Bulk Upload ---")
    csv_data = """email,full_name,roll_number,department,batch_year,cgpa,active_backlogs,skills,resume_text
csv_student1@placematch.edu,CSV Student 1,ROLL_CSV_01,Computer Science,2025,8.1,0,"Python;SQL","CSV Resume text"
csv_student2@placematch.edu,CSV Student 2,ROLL_CSV_02,Information Tech,2025,7.9,0,"Java;React","CSV Resume text"
"""
    csv_files = {"file": ("students.csv", csv_data, "text/csv")}
    csv_res = client.post(f"/api/v1/sections/{section_id}/bulk-upload", files=csv_files, headers=admin_headers)
    print(f"CSV Bulk Upload Status: {csv_res.status_code}")
    print(f"CSV Result: {csv_res.json()}")
    assert csv_res.status_code == 200
    assert csv_res.json()["successful_inserts"] == 2
    print("[SUCCESS] CSV Bulk Upload succeeded!")

    # -------------------------------------------------------------------------
    # TEST 4: Excel (.xlsx) Bulk Upload
    # -------------------------------------------------------------------------
    print("\n--- TEST 4: Excel (.xlsx) Bulk Upload ---")
    excel_df = pd.DataFrame([
        {
            "Email": "excel_student1@placematch.edu",
            "Full Name": "Excel Student 1",
            "Roll Number": "ROLL_XLSX_01",
            "Department": "Cybersecurity",
            "Batch Year": 2025,
            "CGPA": 9.1,
            "Active Backlogs": 0,
            "Skills": "Python; Docker; Kubernetes",
            "Resume Text": "Excel resume text for candidate 1"
        },
        {
            "Email": "excel_student2@placematch.edu",
            "Full Name": "Excel Student 2",
            "Roll Number": "ROLL_XLSX_02",
            "Department": "Data Science",
            "Batch Year": 2025,
            "CGPA": 8.5,
            "Active Backlogs": 0,
            "Skills": "Python; Pandas; PyTorch",
            "Resume Text": "Excel resume text for candidate 2"
        }
    ])
    excel_buffer = io.BytesIO()
    excel_df.to_excel(excel_buffer, index=False, engine="openpyxl")
    excel_bytes = excel_buffer.getvalue()

    excel_files = {"file": ("students.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    excel_res = client.post(f"/api/v1/sections/{section_id}/bulk-upload", files=excel_files, headers=admin_headers)
    print(f"Excel Bulk Upload Status: {excel_res.status_code}")
    print(f"Excel Result: {excel_res.json()}")
    assert excel_res.status_code == 200
    assert excel_res.json()["successful_inserts"] == 2
    print("[SUCCESS] Excel (.xlsx) Bulk Upload succeeded!")

    # -------------------------------------------------------------------------
    # TEST 5: Word (.docx) Table Bulk Upload
    # -------------------------------------------------------------------------
    print("\n--- TEST 5: Word (.docx) Table Bulk Upload ---")
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=9)
    hdr_cells = table.rows[0].cells
    headers = ["Email", "Full Name", "Roll Number", "Department", "Batch Year", "CGPA", "Active Backlogs", "Skills", "Resume Text"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h

    word_data = [
        ["word_student1@placematch.edu", "Word Student 1", "ROLL_DOCX_01", "Computer Science", "2025", "8.9", "0", "Java; Spring; MySQL", "Word doc resume text 1"],
        ["word_student2@placematch.edu", "Word Student 2", "ROLL_DOCX_02", "Information Tech", "2025", "9.4", "0", "React; Node; AWS", "Word doc resume text 2"],
    ]
    for row in word_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    docx_files = {"file": ("students.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    docx_res = client.post(f"/api/v1/sections/{section_id}/bulk-upload", files=docx_files, headers=admin_headers)
    print(f"Word (.docx) Bulk Upload Status: {docx_res.status_code}")
    print(f"Word Result: {docx_res.json()}")
    assert docx_res.status_code == 200
    assert docx_res.json()["successful_inserts"] == 2
    print("[SUCCESS] Word (.docx) Table Bulk Upload succeeded!")

    # -------------------------------------------------------------------------
    # TEST 6: Word (.docx) Invalid Format Error Handling
    # -------------------------------------------------------------------------
    print("\n--- TEST 6: Word (.docx) Invalid Format Error Handling ---")
    bad_doc = docx.Document()
    bad_doc.add_paragraph("This is just plain text paragraph with no student table or CSV format.")
    bad_buffer = io.BytesIO()
    bad_doc.save(bad_buffer)

    bad_files = {"file": ("invalid.docx", bad_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    bad_res = client.post(f"/api/v1/sections/{section_id}/bulk-upload", files=bad_files, headers=admin_headers)
    print(f"Invalid Word Upload Status: {bad_res.status_code}")
    print(f"Invalid Word Response: {bad_res.json()}")
    assert bad_res.status_code == 400
    assert "Could not find a valid student data table" in bad_res.json()["detail"]
    print("[SUCCESS] Word (.docx) invalid format returned clear HTTP 400 error message!")

    print("\n[SUCCESS] ALL MULTI-FORMAT UPLOAD TESTS PASSED PERFECTLY!")

if __name__ == "__main__":
    run_tests()
