from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Query, Form
from sqlalchemy.orm import Session
from typing import List, Optional, Dict, Any
from decimal import Decimal
import csv
import io

from backend.app.core.database import get_db
from backend.app.core.security import require_role, get_password_hash
from backend.app.models.user import User, UserRole
from backend.app.models.student import Student
from backend.app.models.section import Section
from backend.app.models.audit_log import AuditLog
from backend.app.schemas.section import SectionCreate, SectionResponse
from backend.app.schemas.student import StudentResponse, StudentUpdate, StudentBulkUploadResult

router = APIRouter(prefix="/sections", tags=["Sections"])

_admin_only = require_role([UserRole.ADMIN])
_admin_or_recruiter = require_role([UserRole.ADMIN, UserRole.RECRUITER])


def _section_or_404(section_id: int, db: Session) -> Section:
    section = db.query(Section).filter(Section.id == section_id).first()
    if not section:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Section not found.")
    return section


def _to_section_response(section: Section, db: Session) -> SectionResponse:
    count = db.query(Student).filter(Student.section_id == section.id).count()
    return SectionResponse(
        id=section.id,
        name=section.name,
        created_by=section.created_by,
        created_at=section.created_at,
        student_count=count,
    )


def _normalize_header_key(key: Any) -> str:
    if key is None:
        return ""
    k = str(key).strip().lower().replace(" ", "_").replace("-", "_")
    mapping = {
        "name": "full_name",
        "student_name": "full_name",
        "email_address": "email",
        "roll_no": "roll_number",
        "rollno": "roll_number",
        "dept": "department",
        "branch": "department",
        "batch": "batch_year",
        "gpa": "cgpa",
        "backlogs": "active_backlogs",
        "resume": "resume_text",
    }
    return mapping.get(k, k)


def _extract_text_from_file(file_name: str, content: bytes) -> str:
    """Extract text from uploaded resume file (supports PDF, TXT, or plain text)."""
    fn_lower = file_name.lower()
    if fn_lower.endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            extracted_pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(extracted_pages).strip()
            if text:
                return text
        except Exception as e:
            print(f"Warning: Failed pypdf extraction for {file_name}: {e}")

    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


def parse_bulk_student_file(file_name: str, contents: bytes) -> List[Dict[str, Any]]:
    """Parse bulk student data from CSV (.csv, .txt), Excel (.xlsx, .xls), or Word (.docx) files."""
    fn_lower = file_name.lower()

    if fn_lower.endswith(".csv") or fn_lower.endswith(".txt"):
        try:
            decoded = contents.decode("utf-8")
        except UnicodeDecodeError:
            decoded = contents.decode("latin-1", errors="ignore")

        reader = csv.DictReader(io.StringIO(decoded))
        rows = []
        for r in reader:
            normalized_row = {_normalize_header_key(k): v for k, v in r.items() if k is not None}
            rows.append(normalized_row)
        return rows

    elif fn_lower.endswith(".xlsx") or fn_lower.endswith(".xls"):
        try:
            import pandas as pd
            df = pd.read_excel(io.BytesIO(contents))
            df = df.fillna("")
            raw_dict = df.to_dict(orient="records")
            rows = []
            for r in raw_dict:
                normalized_row = {_normalize_header_key(k): str(v).strip() if v != "" else "" for k, v in r.items()}
                rows.append(normalized_row)
            return rows
        except Exception as err:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse Excel file (.xlsx/.xls): {str(err)}"
            )

    elif fn_lower.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(contents))
            rows = []

            # 1. Primary case: Table-based docx
            if doc.tables and len(doc.tables[0].rows) > 1:
                table = doc.tables[0]
                header_cells = [cell.text.strip() for cell in table.rows[0].cells]
                headers = [_normalize_header_key(h) for h in header_cells]

                for row in table.rows[1:]:
                    row_data = {}
                    for col_idx, cell in enumerate(row.cells):
                        if col_idx < len(headers):
                            row_data[headers[col_idx]] = cell.text.strip()
                    if any(row_data.values()):
                        rows.append(row_data)

            # 2. Fallback case: Paragraph-based key-value or CSV formatted text
            if not rows:
                current_record = {}
                for para in doc.paragraphs:
                    line = para.text.strip()
                    if not line:
                        if current_record and ("email" in current_record or "roll_number" in current_record):
                            rows.append(current_record)
                            current_record = {}
                        continue

                    if ":" in line:
                        parts = line.split(":", 1)
                        k = _normalize_header_key(parts[0].strip())
                        v = parts[1].strip()
                        current_record[k] = v
                    elif "," in line:
                        cols = [c.strip() for c in line.split(",")]
                        if len(cols) >= 3:
                            rows.append({
                                "email": cols[0],
                                "full_name": cols[1],
                                "roll_number": cols[2],
                                "department": cols[3] if len(cols) > 3 else "Computer Science",
                                "cgpa": cols[4] if len(cols) > 4 else "7.0",
                            })

                if current_record and ("email" in current_record or "roll_number" in current_record):
                    rows.append(current_record)

            if not rows:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Could not find a valid student data table or formatted student records in the Word document (.docx)."
                )
            return rows

        except HTTPException:
            raise
        except Exception as err:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse Word document (.docx): {str(err)}"
            )

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a .csv, .xlsx, .xls, or .docx file."
        )


def _run_and_store_section_ai_analysis(section_id: int, db: Session):
    """Automatically execute and store section-scoped Isolation Forest anomaly score
    and SHAP analysis for every student in the section into AuditLog."""
    peer_students = db.query(Student).filter(Student.section_id == section_id).all()
    if not peer_students:
        return

    from backend.app.services.anomaly_service import get_anomaly_detection_service
    from backend.app.services.shap_service import get_shap_comparison_service

    anomaly_svc = get_anomaly_detection_service()
    shap_svc = get_shap_comparison_service()
    section = db.query(Section).filter(Section.id == section_id).first()
    section_name = section.name if section else f"Section #{section_id}"

    for s in peer_students:
        peers_for_s = [p for p in peer_students if p.id != s.id]
        
        # 1. Scoped Anomaly score
        scoped_result = anomaly_svc.evaluate_student_anomaly(
            resume_text=s.raw_resume_text or "",
            skills=s.skills or [],
            projects=s.projects or [],
            peer_students=peers_for_s,
        )
        global_result = anomaly_svc.evaluate_student_anomaly_global(
            resume_text=s.raw_resume_text or "",
            skills=s.skills or [],
            projects=s.projects or [],
        )
        delta = round(scoped_result["anomaly_score"] - global_result["anomaly_score"], 3)
        anomaly_payload = {
            "student_id": s.id,
            "roll_number": s.roll_number,
            "student_name": s.full_name,
            "scope": f"section:{section_id}",
            "section_name": section_name,
            "peer_count": len(peers_for_s),
            "global_anomaly_score": global_result["anomaly_score"],
            "scoped_anomaly_score": scoped_result["anomaly_score"],
            "score_delta": delta,
            "delta_direction": "more_anomalous_in_section" if delta > 0 else ("less_anomalous_in_section" if delta < 0 else "identical"),
            "anomaly_audit": scoped_result,
            "global_audit": global_result,
        }

        # 2. Scoped SHAP analysis
        shap_analysis = shap_svc.compute_shap_comparison(
            resume_text=s.raw_resume_text or "",
            skills=s.skills or [],
            projects=s.projects or [],
            target_domain=s.predicted_domain,
            peer_students=peers_for_s,
        )
        shap_payload = {
            "student_id": s.id,
            "roll_number": s.roll_number,
            "student_name": s.full_name,
            "scope": f"section:{section_id}",
            "section_name": section_name,
            "peer_count": len(peers_for_s),
            "shap_analysis": shap_analysis,
        }

        audit_log = db.query(AuditLog).filter(
            AuditLog.student_id == s.id,
            AuditLog.section_id == section_id
        ).first()
        if not audit_log:
            audit_log = AuditLog(student_id=s.id, section_id=section_id)
            db.add(audit_log)
        audit_log.anomaly_result = anomaly_payload
        audit_log.shap_result = shap_payload

    db.commit()


# ---------------------------------------------------------------------------
# Section CRUD
# ---------------------------------------------------------------------------

@router.post("/", response_model=SectionResponse, status_code=status.HTTP_201_CREATED)
def create_section(
    section_in: SectionCreate,
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Create a new student section (e.g. "IT-A", "CSE-B"). Admin only."""
    existing = db.query(Section).filter(Section.name == section_in.name).first()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Section named '{section_in.name}' already exists."
        )
    section = Section(name=section_in.name, created_by=current_user.id)
    db.add(section)
    db.commit()
    db.refresh(section)
    return _to_section_response(section, db)


@router.get("/", response_model=List[SectionResponse])
def list_sections(
    current_user: User = Depends(_admin_or_recruiter),
    db: Session = Depends(get_db),
):
    """List all sections (Admin and Recruiter read-only)."""
    sections = db.query(Section).order_by(Section.name).all()
    return [_to_section_response(s, db) for s in sections]


@router.get("/{section_id}", response_model=SectionResponse)
def get_section(
    section_id: int,
    current_user: User = Depends(_admin_or_recruiter),
    db: Session = Depends(get_db),
):
    """Get details of a section (Admin and Recruiter read-only)."""
    section = _section_or_404(section_id, db)
    return _to_section_response(section, db)


@router.delete("/{section_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_section(
    section_id: int,
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Delete a section. Admin only."""
    section = _section_or_404(section_id, db)
    db.delete(section)
    db.commit()


# ---------------------------------------------------------------------------
# Student membership management
# ---------------------------------------------------------------------------

@router.get("/{section_id}/students", response_model=List[StudentResponse])
def list_section_students(
    section_id: int,
    skip: int = Query(0, ge=0),
    limit: int = Query(200, ge=1, le=500),
    search: Optional[str] = Query(None),
    current_user: User = Depends(_admin_or_recruiter),
    db: Session = Depends(get_db),
):
    """Return all students in the given section (Admin & Recruiter read-only)."""
    _section_or_404(section_id, db)
    query = (
        db.query(Student)
        .join(User, Student.user_id == User.id)
        .filter(Student.section_id == section_id)
    )
    if search:
        from sqlalchemy import or_
        query = query.filter(
            or_(
                Student.roll_number.ilike(f"%{search}%"),
                User.full_name.ilike(f"%{search}%"),
                User.email.ilike(f"%{search}%"),
            )
        )
    return query.offset(skip).limit(limit).all()


@router.post("/{section_id}/add-student/{student_id}", response_model=StudentResponse)
def add_student_to_section(
    section_id: int,
    student_id: int,
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Assign an existing student to this section. Admin only."""
    _section_or_404(section_id, db)
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student not found.")
    student.section_id = section_id
    db.commit()
    db.refresh(student)

    _run_and_store_section_ai_analysis(section_id, db)
    return student


@router.delete("/{section_id}/students/{student_id}", status_code=status.HTTP_200_OK)
def remove_student_from_section(
    section_id: int,
    student_id: int,
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Remove a student from the section (sets section_id to NULL). Admin only."""
    _section_or_404(section_id, db)
    student = db.query(Student).filter(
        Student.id == student_id, Student.section_id == section_id
    ).first()
    if not student:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student not found in this section."
        )
    student.section_id = None
    db.query(AuditLog).filter(
        AuditLog.student_id == student_id,
        AuditLog.section_id == section_id
    ).delete()
    db.commit()
    return {"message": f"Student {student_id} removed from section {section_id}."}


@router.put("/{section_id}/students/{student_id}", response_model=StudentResponse)
def update_section_student(
    section_id: int,
    student_id: int,
    student_in: StudentUpdate,
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Edit a student's record within a section. Admin only. Invalidates cached AuditLog."""
    _section_or_404(section_id, db)
    student = db.query(Student).filter(
        Student.id == student_id, Student.section_id == section_id
    ).first()
    if not student:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Student not found in this section."
        )

    if student_in.full_name is not None and student.user:
        student.user.full_name = student_in.full_name
    if student_in.roll_number is not None:
        student.roll_number = student_in.roll_number
    if student_in.department is not None:
        student.department = student_in.department
    if student_in.batch_year is not None:
        student.batch_year = student_in.batch_year
    if student_in.cgpa is not None:
        student.cgpa = student_in.cgpa
    if student_in.active_backlogs is not None:
        student.active_backlogs = student_in.active_backlogs
    if student_in.history_backlogs is not None:
        student.history_backlogs = student_in.history_backlogs
    if student_in.phone is not None:
        student.phone = student_in.phone
    if student_in.skills is not None:
        student.skills = student_in.skills
    if student_in.raw_resume_text is not None:
        student.raw_resume_text = student_in.raw_resume_text
    if student_in.predicted_domain is not None:
        student.predicted_domain = student_in.predicted_domain

    if student_in.raw_resume_text is not None or student_in.skills is not None:
        try:
            from backend.app.services.classifier_service import get_classifier_service
            clf = get_classifier_service()
            result = clf.predict_and_explain(
                resume_text=student.raw_resume_text or "",
                skills=student.skills or [],
                projects=student.projects or [],
            )
            student.predicted_domain = result["predicted_domain"]
            student.domain_confidence = result["confidence"]
        except Exception:
            pass

    db.query(AuditLog).filter(
        AuditLog.student_id == student_id,
        AuditLog.section_id == section_id
    ).delete()

    db.commit()
    db.refresh(student)

    _run_and_store_section_ai_analysis(section_id, db)
    return student


# ---------------------------------------------------------------------------
# Multi-format Section-scoped Bulk Upload (CSV, Excel, Word)
# ---------------------------------------------------------------------------

@router.post("/{section_id}/bulk-upload", response_model=StudentBulkUploadResult)
async def section_bulk_upload_file(
    section_id: int,
    file: UploadFile = File(...),
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Upload a file (CSV, Excel .xlsx/.xls, Word .docx) of students directly into a section. Admin only.
    Automatically executes and stores section-scoped Isolation Forest + SHAP analysis."""
    _section_or_404(section_id, db)

    contents = await file.read()
    parsed_rows = parse_bulk_student_file(file.filename, contents)

    total_parsed = 0
    successful_inserts = 0
    failed_rows = []
    default_hash = get_password_hash("password123")

    for row_idx, row in enumerate(parsed_rows, start=2):
        total_parsed += 1
        try:
            email = str(row.get("email", "")).strip()
            full_name = str(row.get("full_name", "")).strip()
            roll_number = str(row.get("roll_number", "")).strip()
            department = str(row.get("department", "Computer Science")).strip() or "Computer Science"
            batch_year = int(float(row.get("batch_year", 2025) or 2025))
            cgpa_val = row.get("cgpa", 7.00)
            cgpa = Decimal(str(cgpa_val if cgpa_val != "" else 7.00))
            active_backlogs = int(float(row.get("active_backlogs", 0) or 0))

            skills_raw = str(row.get("skills", ""))
            if ";" in skills_raw:
                skills = [s.strip() for s in skills_raw.split(";") if s.strip()]
            elif "," in skills_raw:
                skills = [s.strip() for s in skills_raw.split(",") if s.strip()]
            else:
                skills = [skills_raw.strip()] if skills_raw.strip() else []

            resume_text = str(row.get("resume_text", f"Resume summary for {full_name}"))

            if not email or not roll_number or not full_name:
                failed_rows.append({
                    "row": row_idx,
                    "reason": "Missing required fields (email, full_name, or roll_number)"
                })
                continue

            existing_user = db.query(User).filter(User.email == email).first()
            if existing_user:
                existing_student = db.query(Student).filter(Student.user_id == existing_user.id).first()
                if existing_student:
                    existing_student.section_id = section_id
                    successful_inserts += 1
                    continue
                else:
                    failed_rows.append({"row": row_idx, "email": email, "reason": "Email already exists"})
                    continue

            existing_roll = db.query(Student).filter(Student.roll_number == roll_number).first()
            if existing_roll:
                existing_roll.section_id = section_id
                successful_inserts += 1
                continue

            new_user = User(
                email=email,
                hashed_password=default_hash,
                role=UserRole.STUDENT,
                full_name=full_name,
                is_active=True,
            )
            db.add(new_user)
            db.flush()

            new_student = Student(
                user_id=new_user.id,
                roll_number=roll_number,
                department=department,
                batch_year=batch_year,
                cgpa=cgpa,
                active_backlogs=active_backlogs,
                history_backlogs=active_backlogs,
                skills=skills,
                raw_resume_text=resume_text,
                predicted_domain=row.get("predicted_domain", "Software Development"),
                domain_confidence=0.85,
                section_id=section_id,
            )
            db.add(new_student)
            successful_inserts += 1

        except Exception as err:
            failed_rows.append({"row": row_idx, "reason": str(err)})

    db.commit()

    try:
        _run_and_store_section_ai_analysis(section_id, db)
    except Exception as e:
        print(f"Warning: Failed to run automatic AI analysis post bulk-upload: {e}")

    return {
        "total_parsed": total_parsed,
        "successful_inserts": successful_inserts,
        "failed_rows": failed_rows,
    }


# ---------------------------------------------------------------------------
# Section-scoped single resume upload with Global Data Consistency Check
# ---------------------------------------------------------------------------

@router.post("/{section_id}/upload-resume", response_model=StudentResponse)
async def section_upload_resume(
    section_id: int,
    email: str = Query(..., description="Student email"),
    full_name: str = Query(..., description="Student full name"),
    roll_number: str = Query(..., description="Student roll number"),
    department: str = Query("Computer Science"),
    cgpa: float = Query(7.5),
    active_backlogs: int = Query(0),
    skills_csv: Optional[str] = Query(None, description="Comma-separated skills"),
    file: Optional[UploadFile] = File(None),
    raw_text: Optional[str] = Query(None, description="Raw text if file not uploaded"),
    confirm_overwrite: bool = Query(False, description="Explicitly confirm overwriting conflicting record"),
    resolution: Optional[str] = Query(None, description="'overwrite' or 'keep_existing'"),
    current_user: User = Depends(_admin_only),
    db: Session = Depends(get_db),
):
    """Upload a single resume/student into a section with global consistency check across ALL sections and unassigned students."""
    _section_or_404(section_id, db)

    extracted_text = ""
    if file:
        content = await file.read()
        extracted_text = _extract_text_from_file(file.filename, content)
    elif raw_text:
        extracted_text = raw_text
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Either a resume file or raw_text must be provided."
        )

    skills_list = [s.strip() for s in skills_csv.split(",") if s.strip()] if skills_csv else []

    existing_user = db.query(User).filter(User.email == email).first()
    existing_student = None
    if existing_user:
        existing_student = db.query(Student).filter(Student.user_id == existing_user.id).first()
    if not existing_student:
        existing_student = db.query(Student).filter(Student.roll_number == roll_number).first()
        if existing_student and not existing_user:
            existing_user = existing_student.user

    if existing_student:
        conflicting_fields = []

        if existing_user and existing_user.full_name != full_name:
            conflicting_fields.append({
                "field": "full_name",
                "existing_value": existing_user.full_name,
                "incoming_value": full_name
            })
        if existing_student.roll_number != roll_number:
            conflicting_fields.append({
                "field": "roll_number",
                "existing_value": existing_student.roll_number,
                "incoming_value": roll_number
            })
        if existing_user and existing_user.email != email:
            conflicting_fields.append({
                "field": "email",
                "existing_value": existing_user.email,
                "incoming_value": email
            })
        if existing_student.department != department:
            conflicting_fields.append({
                "field": "department",
                "existing_value": existing_student.department,
                "incoming_value": department
            })
        existing_cgpa = float(existing_student.cgpa) if existing_student.cgpa is not None else None
        if existing_cgpa != float(cgpa):
            conflicting_fields.append({
                "field": "cgpa",
                "existing_value": existing_cgpa,
                "incoming_value": float(cgpa)
            })
        if existing_student.active_backlogs != active_backlogs:
            conflicting_fields.append({
                "field": "active_backlogs",
                "existing_value": existing_student.active_backlogs,
                "incoming_value": active_backlogs
            })
        existing_skills = existing_student.skills or []
        if sorted(existing_skills) != sorted(skills_list):
            conflicting_fields.append({
                "field": "skills",
                "existing_value": existing_skills,
                "incoming_value": skills_list
            })
        if (existing_student.raw_resume_text or "").strip() != (extracted_text or "").strip():
            conflicting_fields.append({
                "field": "raw_resume_text",
                "existing_value": existing_student.raw_resume_text,
                "incoming_value": extracted_text
            })

        if conflicting_fields and not confirm_overwrite and resolution not in ["overwrite", "keep_existing"]:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail={
                    "message": "Data consistency conflict detected: Student record already exists with differing values.",
                    "status": "conflict",
                    "existing_student_id": existing_student.id,
                    "conflicting_fields": conflicting_fields,
                    "existing_record": {
                        "full_name": existing_user.full_name if existing_user else "",
                        "email": existing_user.email if existing_user else "",
                        "roll_number": existing_student.roll_number,
                        "department": existing_student.department,
                        "cgpa": float(existing_student.cgpa) if existing_student.cgpa is not None else None,
                        "active_backlogs": existing_student.active_backlogs,
                        "skills": existing_student.skills or [],
                        "raw_resume_text": existing_student.raw_resume_text,
                        "section_id": existing_student.section_id
                    },
                    "incoming_record": {
                        "full_name": full_name,
                        "email": email,
                        "roll_number": roll_number,
                        "department": department,
                        "cgpa": float(cgpa),
                        "active_backlogs": active_backlogs,
                        "skills": skills_list,
                        "raw_resume_text": extracted_text,
                        "section_id": section_id
                    }
                }
            )

        if resolution == "keep_existing":
            existing_student.section_id = section_id
            db.commit()
            db.refresh(existing_student)
            _run_and_store_section_ai_analysis(section_id, db)
            return existing_student

        existing_student.department = department
        existing_student.cgpa = Decimal(str(cgpa))
        existing_student.active_backlogs = active_backlogs
        if skills_list:
            existing_student.skills = skills_list
        existing_student.raw_resume_text = extracted_text
        if existing_user:
            existing_user.full_name = full_name
        existing_student.section_id = section_id

        try:
            from backend.app.services.classifier_service import get_classifier_service
            clf = get_classifier_service()
            res = clf.predict_and_explain(
                resume_text=existing_student.raw_resume_text or "",
                skills=existing_student.skills or [],
                projects=existing_student.projects or [],
            )
            existing_student.predicted_domain = res["predicted_domain"]
            existing_student.domain_confidence = res["confidence"]
        except Exception:
            existing_student.predicted_domain = "Software Development"
            existing_student.domain_confidence = 0.85

        db.commit()
        db.refresh(existing_student)

        _run_and_store_section_ai_analysis(section_id, db)
        return existing_student

    # User exists (e.g. registered via auth) but has no Student profile yet — reuse the User
    if existing_user:
        user = existing_user
        if user.full_name != full_name:
            user.full_name = full_name
    else:
        default_hash = get_password_hash("password123")
        user = User(
            email=email,
            hashed_password=default_hash,
            role=UserRole.STUDENT,
            full_name=full_name,
            is_active=True,
        )
        db.add(user)
        db.flush()

    student = Student(
        user_id=user.id,
        roll_number=roll_number,
        department=department,
        batch_year=2025,
        cgpa=Decimal(str(cgpa)),
        active_backlogs=active_backlogs,
        skills=skills_list,
        raw_resume_text=extracted_text,
        section_id=section_id,
    )
    db.add(student)
    db.flush()

    try:
        from backend.app.services.classifier_service import get_classifier_service
        clf = get_classifier_service()
        res = clf.predict_and_explain(
            resume_text=student.raw_resume_text or "",
            skills=student.skills or [],
            projects=student.projects or [],
        )
        student.predicted_domain = res["predicted_domain"]
        student.domain_confidence = res["confidence"]
    except Exception:
        student.predicted_domain = "Software Development"
        student.domain_confidence = 0.85

    db.commit()
    db.refresh(student)

    _run_and_store_section_ai_analysis(section_id, db)
    return student

